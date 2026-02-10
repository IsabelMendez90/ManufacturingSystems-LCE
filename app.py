# LCE + 5S Manufacturing System & Supply Chain Decision Support — AI Agent
# Authors: Dr. J. Isabel Méndez & Dr. Arturo Molina
# Notes:
# - Bounded AI AGENT (planner → tools → reflector) that USES an LLM.
# - LCE Stage selection and Stage Views (Function/Org/Info/Resource/Performance).

import streamlit as st
import pandas as pd
import plotly.express as px
from openai import OpenAI
from fpdf import FPDF
from io import BytesIO
from datetime import datetime
import json, re
import hashlib
import math
import random

# ------------------------ App + API setup ------------------------
st.set_page_config(page_title="LCE + 5S", layout="wide")
st.title("LCE + 5S Manufacturing System & Supply Chain Decision Support (AI Agent)")
st.markdown("Developed by: **Dr. J. Isabel Méndez** & **Dr. Arturo Molina**")

st.markdown("""
    <style>
    /* Primary style for ALL st.button instances */
    div.stButton > button:first-child {
        background-color: #00785D !important;
        color: white !important;
        font-size: 1.3rem !important;
        font-weight: bold !important;
        padding: 22px 0 !important;
        width: 100% !important;
        border-radius: 15px !important;
        margin-top: 18px !important;
        margin-bottom: 18px !important;
        border: 0 !important;
    }
    div.stButton > button:first-child:hover {
        background-color: #009874 !important;
    }

    /* Optional: make the Download button match */
    div.stDownloadButton > button {
        background-color: #00785D !important;
        color: white !important;
        font-size: 1.1rem !important;
        font-weight: 700 !important;
        padding: 14px 20px !important;
        border-radius: 12px !important;
        border: 0 !important;
    }
    div.stDownloadButton > button:hover {
        background-color: #009874 !important;
    }
    </style>
""", unsafe_allow_html=True)

API_KEY = st.secrets["OPENROUTER_API_KEY"]
client = OpenAI(base_url="https://openrouter.ai/api/v1", api_key=API_KEY)

TXT_LIMIT = 12000  # keep LLM prompts bounded

# ------------------------ Evaluation metrics (optional) ------------------------
EVAL_SOURCES = ["Measured", "Simulated", "Estimated", "Illustrative (example only)"]
EVAL_METRICS = [
    {"category": "Social", "metric": "Training hours per employee", "unit": "hours/employee/year", "direction": "higher", "driver": "Social", "min": 5, "max": 60},
    {"category": "Sustainable", "metric": "Energy per unit", "unit": "kWh/unit", "direction": "lower", "driver": "Sustainable", "min": 10, "max": 100},
    {"category": "Sustainable", "metric": "CO2e per unit", "unit": "kg CO2e/unit", "direction": "lower", "driver": "Sustainable", "min": 1, "max": 20},
    {"category": "Sensing", "metric": "Sensor coverage", "unit": "% of critical assets", "direction": "higher", "driver": "Sensing", "min": 10, "max": 100},
    {"category": "Smart", "metric": "OEE", "unit": "%", "direction": "higher", "driver": "Smart", "min": 40, "max": 90},
    {"category": "Safe", "metric": "TRIR", "unit": "cases per 200k hours", "direction": "lower", "driver": "Safe", "min": 0.2, "max": 6.0},
    {"category": "Supply Chain", "metric": "Lead time", "unit": "days", "direction": "lower", "driver": "SC", "min": 5, "max": 40},
    {"category": "Supply Chain", "metric": "WIP", "unit": "units", "direction": "lower", "driver": "SC", "min": 50, "max": 2000},
]

# ------------------------ Helpers: file text extraction (best-effort) ------------------------
def try_extract_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    data = uploaded_file.read()
    # Simple text/markdown/CSV/log
    if name.endswith((".txt", ".md", ".csv", ".log")):
        try:
            return data.decode("utf-8", errors="ignore")
        except Exception:
            return data.decode("latin-1", errors="ignore")
    # DOCX (optional)
    if name.endswith(".docx"):
        try:
            import docx
            f = BytesIO(data)
            doc = docx.Document(f)
            return "\n".join([p.text for p in doc.paragraphs])
        except Exception:
            return ""
    # PDF (very light, if PyMuPDF present)
    if name.endswith(".pdf"):
        try:
            import fitz  # PyMuPDF
            pdf = fitz.open(stream=data, filetype="pdf")
            text = []
            for i in range(min(20, pdf.page_count)):
                text.append(pdf.load_page(i).get_text())
            return "\n".join(text)
        except Exception:
            return ""
    # Fallback
    return ""

# ------------------------ LLM wrapper ------------------------
def llm(msgs, temperature=0.0, seed=42, model="mistralai/mistral-7b-instruct"):
    resp = client.chat.completions.create(model=model, temperature=temperature, seed=seed, messages=msgs)
    return resp.choices[0].message.content or ""

# ------------------------ 5S Taxonomy (condensed) ------------------------
five_s_taxonomy = {
    "Social":[
        {"desc":"No integration of social factors; minimal worker well-being consideration.","tech":[]},
        {"desc":"Compliance with basic labor laws and ergonomic guidelines.","tech":["Ergonomic checklists","Manual safety reporting"]},
        {"desc":"Workforce development / inclusion programs.","tech":["LMS (Moodle, SuccessFactors)","Engagement apps"]},
        {"desc":"Worker-centric design; real-time well-being feedback.","tech":["Wearables (fatigue)","Sentiment analysis"]},
        {"desc":"Integrated socio-technical systems (co-creation, mental health, community).","tech":["Inclusion dashboards","Psychological safety analytics","Collab platforms"]}
    ],
    "Sustainable":[
        {"desc":"No environmental consideration.","tech":[]},
        {"desc":"Basic compliance; manual resource tracking.","tech":["Spreadsheets","Basic meters"]},
        {"desc":"Energy efficiency; partial recycling.","tech":["ISO 14001 checklists","Env sensors","Recycling SW"]},
        {"desc":"LCA, closed-loop, carbon monitoring.","tech":["OpenLCA","GaBi","Carbon calculators","WMS recycling metrics"]},
        {"desc":"Circular economy; AI optimization/reporting.","tech":["AI energy optimization","Blockchain traceability","ESG analytics"]}
    ],
    "Sensing":[
        {"desc":"Human-only sensing.","tech":["Visual inspection","Manual logging"]},
        {"desc":"Basic sensors; local alarms.","tech":["Thermocouples","Limit switches","PLC indicators"]},
        {"desc":"Multi-sensor + WSN + partial fusion.","tech":["WSN","SCADA","LabVIEW","Basic DAQ"]},
        {"desc":"Embedded + IoT + real-time feedback.","tech":["Azure/AWS IoT","MCUs","OPC-UA/MQTT"]},
        {"desc":"Edge AI + IIoT + adaptive closed-loop.","tech":["Edge AI devices","IIoT platforms","Sensor fusion"]}
    ],
    "Smart":[
        {"desc":"Manual control/decisions.","tech":["Paper logs","Work instructions"]},
        {"desc":"Basic open-loop control.","tech":["Relay logic","Timers","Basic HMI"]},
        {"desc":"Closed-loop automation.","tech":["PID","PLCs","HMI/SCADA"]},
        {"desc":"Advanced automation + predictive + MES.","tech":["MES (Opcenter/Proficy)","Predictive (Maximo/Senseye)"]},
        {"desc":"AI/ANN/IIoT/twins/big data.","tech":["SageMaker/Vertex","Keras/PyTorch","Digital twins"]}
    ],
    "Safe":[
        {"desc":"Minimal safety.","tech":["Safety posters","Visual checks"]},
        {"desc":"Compliance (ISO 45001).","tech":["Safety audits","Risk matrices"]},
        {"desc":"Machine safety integrated.","tech":["Safety PLCs","Light curtains","E-stops","Interlocks"]},
        {"desc":"Smart safety (sensors + anomaly).","tech":["Vibration sensors","AI failure prediction","HSE dashboards"]},
        {"desc":"AI safety management + cyber-physical.","tech":["Safety twins","AI hazard detection","IEC 62443/NIST CSF"]}
    ]
}

# --------------- LCE Actions Taxonomy ---------------
lce_actions_taxonomy = {
    "Product Transfer": [
        "Ideation: Identify product BOM, materials, and quality requirements.",
        "Basic Development: Define manufacturing and quality control specs for cost, volume, and delivery.",
        "Advanced Development: Evaluate/select suppliers based on capabilities, capacity, compliance.",
        "Launching: Manage in-house assembly, document and check quality controls, train workforce.",
        "End-of-Life: Plan for disassembly, recycling, reuse; manage contract closure and reverse logistics."
    ],
    "Technology Transfer": [
        "Ideation: Capture technical specs (geometry, materials, batch size) for new component/family.",
        "Basic Development: Specify/select new manufacturing technology, equipment, and supporting tools.",
        "Advanced Development: Develop process plan, define control documentation, SOPs.",
        "Launching: Set up, test, and ramp-up equipment; optimize production.",
        "End-of-Life: Retire/adapt equipment; document performance and failures; involve digital twins for lifecycle learning."
    ],
    "Facility Design": [
        "Ideation: Specify new product requirements and associated manufacturing processes.",
        "Basic Development: Select manufacturing systems and equipment for facility.",
        "Advanced Development: Design shop-floor layout, define capacity and manufacturing strategy.",
        "Launching: Build/install/ramp; evaluate.",
        "End-of-Life: Audit facility for reuse/decommission/transform."
    ]
}

# --------------- Supply Chain Recommendations (concise hints) ---------------
supply_chain_recommendations = {
    "Product Transfer":[
        "Agile multi-sourcing; supplier MCDM (AHP/TCO); APQP/SPC; decentralized traceability."
    ],
    "Technology Transfer":[
        "Pilot lines; technical-economic benchmarking; joint dev; predictive + MES/ERP; digital twins."
    ],
    "Facility Design":[
        "Local/resilient networks; make-or-buy; green infra; layout/OEE simulation; cyber-by-design."
    ],
}

# ------------------------ 5S regex critic (reflector for coverage) ------------------------
FIVE_S_PATTERNS = {
    "Social":[r"\bergonom", r"\bhuman factors?\b", r"\bLMS\b", r"\btraining\b", r"\bco-creation\b|\bparticipatory\b"],
    "Sustainable":[r"\bLCA\b", r"\bISO\s?14001\b", r"\bcarbon (footprint|accounting|intensity)\b", r"\bcircular economy\b|\brecycl|\breuse\b", r"\brenewable\b|\benergy (audit|efficiency|management)\b"],
    "Sensing":[r"\bSCADA\b", r"\bOPC[-\s]?UA\b|\bMQTT\b", r"\bDAQ\b", r"\bIIoT\b|\bedge", r"\bcondition monitoring\b"],
    "Smart":[r"\bPLC\b|\bPID\b", r"\bMES\b|\bMOM\b", r"\bpredictive maintenance\b", r"\bdigital twin", r"\boptimization\b|\bAPS\b"],
    "Safe":[r"\bISO\s?45001\b", r"\bISO\s?13849\b|\bIEC\s?61508\b|\bIEC\s?62061\b", r"\bFMEA\b|\bHAZOP\b|\bLOPA\b", r"\bLOTO\b", r"\bIEC\s?62443\b|\bNIST\b"]
}
REQUIRED_MATCHES_BY_LEVEL = {0:0, 1:1, 2:2, 3:3, 4:4}

def find_5s_gaps(plan_text: str, target_levels: dict) -> list:
    text = plan_text or ""
    gaps = []
    for dim, target in target_levels.items():
        if target <= 0:  # when current level is 0, do not force evidence
            continue
        pats = FIVE_S_PATTERNS.get(dim, [])
        hits = set(p for p in pats if re.search(p, text, re.IGNORECASE))
        need = REQUIRED_MATCHES_BY_LEVEL.get(target, 1)
        if len(hits) < need:
            sample = [re.sub(r"\\b", "", p) for p in pats[:5]]
            gaps.append(f"{dim}: need ≥{need} concrete references; found {len(hits)}. Add: {', '.join(sample)}")
    return gaps

# ------------------------ Structure + LCE stage verifiers ------------------------
REQUIRED_SECTIONS = [
    "Supply Chain Configuration & Action Plan",
    "Improvement Opportunities & Risks",
    "Digital/AI Next Steps",
    "Expected 5S Maturity",
]

def find_structure_gaps(plan_text: str) -> list:
    gaps = []
    for head in REQUIRED_SECTIONS:
        body = parse_section(plan_text, head)
        if not body:
            gaps.append(f"Missing or empty section: [{head}]")
    # Ensure Expected 5S lines are explicit
    exp = parse_section(plan_text, "Expected 5S Maturity")
    if exp:
        for dim in ["Social", "Sustainable", "Sensing", "Smart", "Safe"]:
            pat = rf"{dim}\s*[:\-]?\s*(?:L(?:e?vel)?\s*=?\s*)?[0-4]"
            if not re.search(pat, exp, flags=re.IGNORECASE):
                gaps.append(f"Expected 5S line missing: {dim}: <0-4>")
    return gaps

def find_lce_stage_gaps(plan_text: str, selected_stages: list) -> list:
    gaps = []
    if not selected_stages:
        return gaps
    text = plan_text or ""
    stage_keys = [s.split(":")[0].strip() for s in selected_stages]
    for stage in stage_keys:
        if not re.search(rf"\\b{re.escape(stage)}\\b", text, flags=re.IGNORECASE):
            gaps.append(f"LCE stage not mentioned in plan: {stage}")
    return gaps

# ------------------------ Evidence builder ------------------------
def retrieve_domain_evidence(system_type, industry, selected_stages, five_s_levels):
    stages = [s.split(":")[0].strip() for s in selected_stages] or []
    lines = []
    hints = supply_chain_recommendations.get(system_type, [])
    if hints:
        lines.append("SUPPLY-CHAIN HINTS:")
        lines += [f"- {h}" for h in hints]
    if stages:
        lines.append("\nSELECTED LCE STAGES:")
        lines += [f"- {s}" for s in stages]
    lines.append("\n5S TECH HINTS @ CURRENT LEVEL:")
    for dim, lvl in five_s_levels.items():
        techs = five_s_taxonomy[dim][lvl].get("tech", [])
        if techs:
            lines.append(f"- {dim} (L{lvl}): " + "; ".join(techs))
    return "\n".join(lines)

# ------------------------ Plan parsing helpers ------------------------
def parse_section(text, head):
    if not text:
        return ""
    # Match headings like:
    # [Section], ### [Section], **[Section]**, ### **[Section]**
    header_any = re.compile(r"^\s*(?:#+\s*)?(?:\*\*)?\[([^\]]+)\](?:\*\*)?\s*$", re.MULTILINE)
    matches = list(header_any.finditer(text))
    if not matches:
        return ""
    target = head.strip().lower()
    idx = None
    for i, m in enumerate(matches):
        if m.group(1).strip().lower() == target:
            idx = i
            break
    if idx is None:
        return ""
    start = matches[idx].end()
    end = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
    return text[start:end].strip()

def _cache_key(objective, system_type, industry, role, selected_stages, five_s_levels, docs_text):
    doc_hash = hashlib.sha256((docs_text or "").encode("utf-8")).hexdigest()
    payload = {
        "objective": objective,
        "system_type": system_type,
        "industry": industry,
        "role": role,
        "selected_stages": selected_stages or [],
        "five_s_levels": five_s_levels or {},
        "docs_hash": doc_hash,
    }
    raw = json.dumps(payload, sort_keys=True)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()

# ------------------------ Evaluation helpers ------------------------
def _to_float(x):
    if x is None:
        return None
    if isinstance(x, float) and math.isnan(x):
        return None
    if pd.isna(x):
        return None
    try:
        return float(x)
    except Exception:
        return None

def compute_eval_results(df):
    rows = []
    for _, r in df.iterrows():
        b = _to_float(r.get("Baseline"))
        p = _to_float(r.get("Proposed"))
        if b is None or p is None:
            continue
        delta = p - b
        direction = (r.get("Direction") or "").lower()
        # Improvement: positive = better
        if direction == "higher":
            improve = (delta / b * 100.0) if b != 0 else None
        elif direction == "lower":
            improve = ((b - p) / b * 100.0) if b != 0 else None
        else:
            improve = (delta / b * 100.0) if b != 0 else None
        rows.append({
            "Category": r.get("Category"),
            "Metric": r.get("Metric"),
            "Unit": r.get("Unit"),
            "Baseline": round(b, 1),
            "Proposed": round(p, 1),
            "Δ": round(delta, 1),
            "Improvement (%)": None if improve is None else round(improve, 1),
            "Source": r.get("Source"),
        })
    return rows

def _maturity_index(levels: dict, driver: str) -> float:
    if not levels:
        return 0.0
    if driver in levels:
        return float(levels.get(driver, 0)) / 4.0
    if driver == "SC":
        dims = ["Smart", "Sensing", "Sustainable"]
        vals = [float(levels.get(d, 0)) / 4.0 for d in dims]
        return sum(vals) / len(vals)
    if driver == "ALL":
        vals = [float(v) / 4.0 for v in levels.values()]
        return sum(vals) / len(vals) if vals else 0.0
    return 0.0

def _clamp(x, lo, hi):
    return max(lo, min(hi, x))

def synthesize_metrics(current_5s: dict, expected_5s: dict, noise_pct: float = 5.0, seed: int = 42):
    rng = random.Random(seed)
    rows = []
    for m in EVAL_METRICS:
        driver = m.get("driver", "ALL")
        cur_idx = _maturity_index(current_5s, driver)
        exp_idx = _maturity_index(expected_5s or current_5s, driver)
        lo = float(m.get("min", 0))
        hi = float(m.get("max", 1))
        direction = (m.get("direction") or "").lower()
        # Map maturity index to metric range
        if direction == "higher":
            baseline = lo + (hi - lo) * cur_idx
            proposed = lo + (hi - lo) * exp_idx
        else:  # lower is better
            baseline = hi - (hi - lo) * cur_idx
            proposed = hi - (hi - lo) * exp_idx
        # Add small synthetic variability
        def jitter(x):
            if noise_pct <= 0:
                return x
            factor = 1.0 + rng.uniform(-noise_pct, noise_pct) / 100.0
            return x * factor
        baseline = _clamp(jitter(baseline), lo, hi)
        proposed = _clamp(jitter(proposed), lo, hi)
        rows.append({
            "Category": m["category"],
            "Metric": m["metric"],
            "Unit": m["unit"],
            "Direction": m["direction"],
            "Driver": m.get("driver", "ALL"),
            "Min": lo,
            "Max": hi,
            "Baseline": round(baseline, 1),
            "Proposed": round(proposed, 1),
            "Source": "Illustrative (example only)",
        })
    return pd.DataFrame(rows)

def compute_proposed_from_baseline(df, current_5s: dict, expected_5s: dict):
    out = df.copy()
    for i, r in out.iterrows():
        b = _to_float(r.get("Baseline"))
        if b is None:
            continue
        direction = (r.get("Direction") or "").lower()
        driver = r.get("Driver") or r.get("driver") or "ALL"
        lo = float(r.get("Min") or r.get("min") or 0)
        hi = float(r.get("Max") or r.get("max") or 1)
        cur_idx = _maturity_index(current_5s, driver)
        exp_idx = _maturity_index(expected_5s or current_5s, driver)
        delta = max(0.0, exp_idx - cur_idx)
        # Move baseline toward best value by delta fraction
        if direction == "higher":
            proposed = b + (hi - b) * delta
        else:  # lower is better
            proposed = b - (b - lo) * delta
        proposed = _clamp(proposed, lo, hi)
        out.at[i, "Proposed"] = round(proposed, 1)
    return out

def extract_expected_levels(text, fallback):
    out = {}
    if not isinstance(text, str):
        text = ""
    for dim in ["Social", "Sustainable", "Sensing", "Smart", "Safe"]:
        pat = rf"{dim}\s*[:\-]?\s*(?:L(?:e?vel)?\s*=?\s*)?([0-4])"
        m = re.search(pat, text, flags=re.IGNORECASE)
        if not m:
            m = re.search(rf"{dim}.{{0,30}}?([0-4])", text, flags=re.IGNORECASE | re.DOTALL)
        out[dim] = int(m.group(1)) if m else int(fallback.get(dim, 0))
    return out

# ------------------------ Stage Views helpers ------------------------
def build_context_block(objective, system_type, industry, five_s_levels, selected_stages, plan_text=""):
    stages_list = [s.split(":")[0].strip() for s in selected_stages]
    five_s_str = ", ".join([f"{k}: {v}" for k, v in five_s_levels.items()])
    return (
        "PROJECT CONTEXT\n"
        f"- Objective: {objective}\n"
        f"- System Type: {system_type}\n"
        f"- Industry: {industry}\n"
        f"- 5S Levels: {five_s_str}\n"
        f"- Selected LCE Stages: {', '.join(stages_list) if stages_list else 'None'}\n\n"
        "Recent plan (if any):\n"
        f"{plan_text}\n"
    )

def build_stage_views_prompt_json(context_block, selected_stages):
    stages = [s.split(":")[0].strip() for s in selected_stages]
    stages_str = ", ".join(stages) if stages else "None"
    return (
        f"{context_block}\n\n"
        "Return ONLY a JSON object (no prose) with this schema:\n\n"
        "{\n"
        '  "stages": [\n'
        "    {\n"
        f'      "name": "<one of: {stages_str}>",\n'
        '      "Analysis": { "Function": "..." },\n'
        '      "Synthesis": {\n'
        '        "Organization": "...",\n'
        '        "Information":  "...",\n'
        '        "Resource":     "..."\n'
        "      },\n"
        '      "Evaluation": { "Performance": "..." }\n'
        "    }\n"
        "  ]\n"
        "}\n\n"
        "Rules:\n"
        f"- Include EVERY stage exactly once (for: {stages_str}).\n"
        "- Use plain strings (no markdown).\n"
        "- Keep fields concise but concrete.\n"
    )

def parse_stage_views_json(llm_response, selected_stages):
    m = re.search(r"\{.*\}", llm_response, flags=re.DOTALL)
    payload = m.group(0) if m else llm_response
    data = json.loads(payload)
    stage_keys = [s.split(":")[0].strip() for s in selected_stages]
    out = {}
    for item in data.get("stages", []):
        name = (item.get("name") or "").strip()
        if name in stage_keys:
            A = item.get("Analysis", {}) or {}
            S = item.get("Synthesis", {}) or {}
            E = item.get("Evaluation", {}) or {}
            out[name] = {
                "Function":     (A.get("Function")     or "").strip(),
                "Organization": (S.get("Organization") or "").strip(),
                "Information":  (S.get("Information")  or "").strip(),
                "Resource":     (S.get("Resource")     or "").strip(),
                "Performance":  (E.get("Performance")  or "").strip(),
            }
    return out

def parse_stage_views_from_plan(plan_text, selected_stages):
    stage_keys = [s.split(":")[0].strip() for s in selected_stages]
    out = {k: {"Function":"","Organization":"","Information":"","Resource":"","Performance":""} for k in stage_keys}
    union = "|".join(re.escape(k) for k in stage_keys) if stage_keys else ""
    if not union:
        return out
    parts = re.split(rf"(?:^\[({union})\]\s*|^({union})\s*:)", plan_text or "", flags=re.MULTILINE)
    i = 1
    while i < len(parts):
        name = (parts[i] or parts[i+1] or "").strip()
        body = parts[i+2] if i+2 < len(parts) else ""
        if name in out:
            def grab(label):
                m = re.search(rf"{label}\s*:\s*(.*?)(?:(?:\n[A-Z][A-Za-z \-/]+?\s*:)|\Z)", body, flags=re.DOTALL)
                return (m.group(1).strip() if m else "")
            out[name]["Function"]     = grab("Function")
            out[name]["Organization"] = grab("Organization")
            out[name]["Information"]  = grab("Information")
            out[name]["Resource"]     = grab("Resource")
            out[name]["Performance"]  = grab("Performance")
        i += 3
    return out

# ------------------------ NEW: Explain Expected 5S (rationales) ------------------------
def explain_expected_levels(objective, system_type, industry, selected_stages,
                            five_s_levels, expected_5s, plan_text) -> dict:
    """Ask the LLM to justify each S change with 1–2 concrete reasons from the plan."""
    sc  = parse_section(plan_text, "Supply Chain Configuration & Action Plan")
    imp = parse_section(plan_text, "Improvement Opportunities & Risks")
    nxt = parse_section(plan_text, "Digital/AI Next Steps")

    stages_txt = ", ".join([s.split(":")[0].strip() for s in selected_stages]) or "None"
    context = (
        f"OBJECTIVE: {objective}\n"
        f"SYSTEM: {system_type} | INDUSTRY: {industry}\n"
        f"STAGES: {stages_txt}\n"
        "CURRENT 5S: " + ", ".join([f"{k}={v}" for k,v in five_s_levels.items()]) + "\n"
        "EXPECTED 5S: " + ", ".join([f"{k}={v}" for k,v in expected_5s.items()]) + "\n\n"
        "[SUPPLY CHAIN PLAN]\n" + (sc or "—") + "\n\n"
        "[IMPROVEMENT OPPORTUNITIES]\n" + (imp or "—") + "\n\n"
        "[DIGITAL/AI NEXT STEPS]\n" + (nxt or "—")
    )

    shape = (
        "{\n"
        '  "Social": {"current": <int>, "expected": <int>, "why": ["<10-20 words>", "<10-20 words>"]},\n'
        '  "Sustainable": {"current": <int>, "expected": <int>, "why": ["...", "..."]},\n'
        '  "Sensing": {"current": <int>, "expected": <int>, "why": ["...", "..."]},\n'
        '  "Smart": {"current": <int>, "expected": <int>, "why": ["...", "..."]},\n'
        '  "Safe": {"current": <int>, "expected": <int>, "why": ["...", "..."]}\n'
        "}"
    )

    prompt = (
        "Explain concisely WHY each 5S dimension improves from current to expected, using ONLY evidence in the "
        "plan sections above (specific actions, technologies, standards). No generic phrases.\n"
        "If expected equals current, return a brief justification for maintaining the level "
        "(e.g., no explicit actions identified to increase that dimension).\n\n"
        "Return JSON ONLY with this exact shape:\n" + shape
    )

    txt = llm(
        [
            {"role":"system","content":"You produce terse, evidence-based rationales in JSON only."},
            {"role":"user","content": context + "\n\n" + prompt}
        ],
        temperature=0.2, seed=42
    )

    try:
        data = json.loads(re.search(r"\{.*\}", txt, re.DOTALL).group(0))
    except Exception:
        data = {}
    # Fallbacks for missing or empty rationales
    out = {}
    for dim in ["Social", "Sustainable", "Sensing", "Smart", "Safe"]:
        cur = int(five_s_levels.get(dim, 0))
        exp = int(expected_5s.get(dim, cur))
        entry = (data.get(dim, {}) or {})
        reasons = entry.get("why", []) or []
        if not reasons:
            if exp == cur:
                reasons = [f"No explicit actions found to increase {dim}; maintain current level."]
            else:
                reasons = [f"Rationale not returned; please regenerate or refine plan evidence for {dim}."]
        out[dim] = {"current": cur, "expected": exp, "why": reasons[:2]}
    return out

# ------------------------ Agent state ------------------------
class AgentState:
    def __init__(self, objective, system_type, industry, role, five_s_levels, selected_stages, docs_text):
        self.objective = objective
        self.system_type = system_type
        self.industry = industry
        self.role = role
        self.five_s_levels = five_s_levels
        self.selected_stages = selected_stages
        self.docs_text = docs_text or ""
        self.spec_data = {}
        self.observations = []

# ------------------------ Planner (bounded) ------------------------
def plan_with_llm(state: AgentState, evidence: str) -> str:
    lce_txt = "\n".join([f"- {s}" for s in state.selected_stages]) if state.selected_stages else "None"
    s_txt  = "\n".join([f"- {d}: L{v}" for d, v in state.five_s_levels.items()])
    doc_snip = (state.docs_text or "")[:TXT_LIMIT]
    prompt = (
        f"You are an expert in {state.system_type} manufacturing systems for the {state.industry} industry.\n\n"
        "IMPORTANT INSTRUCTIONS:\n"
        "- Scope: manufacturing system typology, LCE stages, 5S (Social, Sustainable, Sensing, Smart, Safe), with supply chain & Industry 5.0.\n"
        "- Never invent data; use the evidence only.\n"
        "- Keep responses concise, precise, and actionable.\n\n"
        f"Company objective: {state.objective}\n"
        f"User role: {state.role}\n\n"
        "Selected LCE stages:\n"
        f"{lce_txt}\n"
        "Current 5S maturity:\n"
        f"{s_txt}\n"
        "Relevant document snippets (if any):\n"
        "<<<DOC>>>\n"
        f"{doc_snip}\n"
        "<<<END-DOC>>>\n\n"
        "Use the evidence strictly:\n"
        f"{evidence}\n\n"
        "Return the following sections (markdown, concise and actionable):\n"
        "[Supply Chain Configuration & Action Plan]\n"
        "[Improvement Opportunities & Risks]\n"
        "[Digital/AI Next Steps]\n"
        "[Expected 5S Maturity]\n"
        "Format rule: each heading must appear on its own line exactly as shown above (no extra markdown around it).\n"
        "Provide EXACTLY these five lines using integers 0–4 (no prose, no letters):\n"
        "Social: <0-4>\n"
        "Sustainable: <0-4>\n"
        "Sensing: <0-4>\n"
        "Smart: <0-4>\n"
        "Safe: <0-4>\n"
    )
    return llm(
        [
            {"role": "system", "content": "You are a digital manufacturing systems expert."},
            {"role": "user", "content": prompt},
        ],
        temperature=0.2,
        seed=42,
    )

# ------------------------ Tools (bounded) ------------------------
class Tool:
    name = "base"
    def run(self, state: AgentState) -> dict:
        return {"tool": self.name, "summary": "noop", "data": {}}

class SpecExtractorLLM(Tool):
    name = "spec_extractor"
    def run(self, state):
        prompt = (
            "Extract structured requirements from the objective and LCE stages.\n\n"
            f"Objective: {state.objective}\n"
            f"LCE stages: {', '.join([s.split(':')[0] for s in state.selected_stages]) or 'None'}\n"
            f"Industry: {state.industry}\n\n"
            "Return JSON ONLY:\n"
            "{\n"
            ' "requirements": ["<short bullets>"],\n'
            ' "constraints": ["<standards, compliance, budget, quality if implied>"],\n'
            ' "success_metrics": ["<if mentioned in text, keep generic (no numbers)>"]\n'
            "}\n"
        )
        txt = llm(
            [
                {"role":"system","content":"You turn free text into compact JSON specs."},
                {"role":"user","content":prompt}
            ],
            temperature=0.1
        )
        try:
            js = json.loads(re.search(r"\{.*\}", txt, re.DOTALL).group(0))
        except Exception:
            js = {"requirements":[],"constraints":[],"success_metrics":[]}
        state.spec_data = js
        return {"tool": self.name, "summary": "Structured specs extracted", "data": js}


class DocIntakeLLM(Tool):
    name = "doc_intake"
    def run(self, state):
        if not state.docs_text.strip():
            return {"tool": self.name, "summary":"No documents provided.", "data":{}}
        text_block = state.docs_text[:TXT_LIMIT]
        prompt = (
            "From the following text, extract constraints, required standards, safety/environmental notes, "
            "and any numeric parameters.\n"
            "Return JSON ONLY with keys: constraints, standards, safety, parameters.\n"
            "TEXT:\n"
            "<<<DOC>>>\n" + text_block + "\n<<<END-DOC>>>"
        )
        txt = llm(
            [
                {"role":"system","content":"You extract constraints from long docs with concise JSON."},
                {"role":"user","content":prompt}
            ],
            temperature=0.2
        )
        try:
            js = json.loads(re.search(r"\{.*\}", txt, re.DOTALL).group(0))
        except Exception:
            js = {"constraints":[],"standards":[],"safety":[],"parameters":{}}
        return {"tool": self.name, "summary":"Doc constraints extracted", "data": js}

TOOL_REGISTRY = {
    "spec_extractor": SpecExtractorLLM(),
    "doc_intake":     DocIntakeLLM(),
}

# ------------------------ Agent runner (planner → tools → reflector) ------------------------
def agent_run(objective, system_type, industry, role, selected_stages, five_s_levels,
              docs_text: str, enable_agent=True, auto_iterate=True, max_steps=3):
    state = AgentState(objective, system_type, industry, role, five_s_levels, selected_stages, docs_text)
    evidence = retrieve_domain_evidence(system_type, industry, selected_stages, five_s_levels)
    plan_text = ""

    steps = 1 if not enable_agent else max_steps
    observations = []
    for step in range(steps):
        # Planner
        plan_text = plan_with_llm(state, evidence)
        # Reflector (coverage critic on the SC plan section)
        sc_text = parse_section(plan_text, "Supply Chain Configuration & Action Plan") or plan_text
        gaps = []
        gaps += find_5s_gaps(sc_text, five_s_levels)
        gaps += find_structure_gaps(plan_text)
        gaps += find_lce_stage_gaps(sc_text, selected_stages)

        # Tools (bounded set)
        run_tools = []
        if step == 0: run_tools.append("spec_extractor")
        if docs_text.strip() and step == 0: run_tools.append("doc_intake")

        observations = []
        for name in run_tools:
            out = TOOL_REGISTRY[name].run(state)
            observations.append(out)

        # Decide whether to iterate
        passed_verify = (len(gaps) == 0)
        if not enable_agent or not auto_iterate or passed_verify:
            evidence += "\n\n[OBSERVATIONS]\n" + "\n".join([f"- {o['tool']}: {o['summary']}" for o in observations])
            break
        critic = []
        if not passed_verify:
            critic.append("VERIFICATION FINDINGS:\n- " + "\n- ".join(gaps[:10]))
        evidence += "\n\n[NEEDS FIX — ADDRESS IN NEXT PASS]\n" + "\n\n".join(critic)

    # Stage Views pass
    context_block = build_context_block(
        objective, system_type, industry, five_s_levels, selected_stages,
        parse_section(plan_text, "Supply Chain Configuration & Action Plan")
    )
    stage_prompt = build_stage_views_prompt_json(context_block, selected_stages)
    stage_views_raw = llm(
        [{"role":"system","content":"You are a digital manufacturing systems expert."},
         {"role":"user","content":stage_prompt}],
        temperature=0.1
    )
    stage_views = {}
    try:
        stage_views = parse_stage_views_json(stage_views_raw, selected_stages)
    except Exception:
        stage_views = {}
    if (not stage_views or not any(v for v in stage_views.values())) and plan_text:
        stage_views = parse_stage_views_from_plan(plan_text, selected_stages)

    return {
        "plan_text": plan_text,
        "stage_views": stage_views,
    }

# ------------------------ UI ------------------------

# 1) Objective / Industry / Role
st.header("1. Define Scenario")
objective = st.text_input(
    "Objective (e.g., launch new product, adopt a new process, expand a facility):",
    value=st.session_state.get("objective", "Design and ramp a flexible small manufacturing cell."),
    key="objective"
)
industry = st.selectbox(
    "Industry:", ["Automotive","Electronics","Medical Devices","Consumer Goods","Other"],
    index=st.session_state.get("industry_idx", 1), key="industry"
)
st.session_state["industry_idx"] = ["Automotive","Electronics","Medical Devices","Consumer Goods","Other"].index(industry)

role_options = ["Design Engineer","Process Engineer","Manufacturing Engineer","Safety Supervisor",
                "Sustainability Manager","Supply Chain Analyst","Manager/Decision Maker","Other"]
role_selected = st.selectbox(
    "Your role:", role_options,
    index=st.session_state.get("role_idx", 2), key="user_role"
)
if role_selected == "Other":
    role_selected = st.text_input("Please specify your role:", value=st.session_state.get("custom_role",""), key="custom_role") or "Other"
st.session_state["role_idx"] = role_options.index("Other") if role_selected=="Other" else role_options.index(role_selected)

# 2) Manufacturing System Type
st.header("2. Select Manufacturing System Type")
system_types = ["Product Transfer","Technology Transfer","Facility Design"]
system_type = st.radio("Manufacturing system type:", system_types, key="system_type")

# 3) Select LCE Stages/Actions
st.header("3. Select Relevant LCE Stages/Actions")
lce_global_keys = ["Ideation","Basic Development","Advanced Development","Launching","End-of-Life"]
if "lce_global_checked" not in st.session_state:
    st.session_state["lce_global_checked"] = [False]*len(lce_global_keys)

lce_actions = lce_actions_taxonomy[system_type]
selected_stages=[]
for i, action in enumerate(lce_actions):
    action_key = action.split(":")[0].strip()
    idx = lce_global_keys.index(action_key)
    checked = st.checkbox(action, value=st.session_state["lce_global_checked"][idx], key=f"lce_{i}")
    st.session_state["lce_global_checked"][idx] = checked
    if checked: selected_stages.append(action)
st.session_state["selected_stages"] = selected_stages

# 4) 5S Maturity
st.header("4. Current 5S Maturity (one per S)")
five_s_levels={}
cols = st.columns(5)
for i, dim in enumerate(["Social","Sustainable","Sensing","Smart","Safe"]):
    with cols[i]:
        options=[f"Level {idx}: {lvl['desc']}" for idx,lvl in enumerate(five_s_taxonomy[dim])]
        sel = st.radio(dim, options, index=0, key=f"{dim}_radio")
        five_s_levels[dim] = options.index(sel)
        techs = five_s_taxonomy[dim][five_s_levels[dim]].get("tech", [])
        if techs: st.caption("Tech hints: " + "; ".join(techs))

# 5) Optional Upload relevant docs
st.header("Optional Upload relevant docs (manuals/specs/SOPs)")
uploads = st.file_uploader(
    "Upload .txt/.md/.csv/.log/.docx/.pdf (max a few MB). Multiple allowed.",
    type=["txt","md","csv","log","docx","pdf"], accept_multiple_files=True
)
docs_text = ""
if uploads:
    pieces=[]
    for f in uploads:
        text = try_extract_text(f)
        if text: pieces.append(f"# {f.name}\n{text}")
    docs_text = "\n\n".join(pieces)

# --------------- Run agent ---------------
enable_agent = True
auto_iterate = True

if st.button("Generate Plan & Recommendations"):
    with st.spinner("Agent planning, executing tools, stage views, and refining..."):
        cache_key = _cache_key(
            objective, system_type, industry, role_selected,
            selected_stages, five_s_levels, docs_text
        )
        st.session_state["current_cache_key"] = cache_key
        if "plan_cache" not in st.session_state:
            st.session_state["plan_cache"] = {}
        if cache_key in st.session_state["plan_cache"]:
            result = st.session_state["plan_cache"][cache_key]
        else:
            result = agent_run(
                objective=objective,
                system_type=system_type,
                industry=industry,
                role=role_selected,
                selected_stages=selected_stages,
                five_s_levels=five_s_levels,
                docs_text=docs_text,
                enable_agent=enable_agent,   # always True
                auto_iterate=auto_iterate,   # always True
                max_steps=3,
            )
            st.session_state["plan_cache"][cache_key] = result
        st.session_state["agent_result"] = result

# ------------------------ Results ------------------------
res = st.session_state.get("agent_result")

if res:
    st.header("Results")
    plan_text = res["plan_text"]

    # 6a) Stage Views — System Type + Selected LCE Stages rendered
    st.subheader("Supply Chain Configuration & Action Plan (by LCE Stage)")
    displayed_stages = st.session_state.get("selected_stages", [])
    stage_views = res.get("stage_views", {}) or {}
    if displayed_stages:
        st.markdown("<ul style='margin-top:0; margin-bottom:0;'>", unsafe_allow_html=True)
        for action in displayed_stages:
            stage, desc = (action.split(":",1)+[""])[:2]
            stage_key = stage.strip(); desc = desc.strip()
            views = stage_views.get(stage_key, {})
            st.markdown(f"<li><b>Stage → {stage_key}</b>: {desc}", unsafe_allow_html=True)
            st.markdown("<ul style='margin-top:0; margin-bottom:0;'>", unsafe_allow_html=True)
            groups = [
                ("Engineering Analysis", [("Function","Function")]),
                ("Engineering Synthesis", [("Organization","Organization"),("Information","Information"),("Resource","Resource")]),
                ("Engineering Evaluation", [("Performance","Performance")]),
            ]
            for heading, items in groups:
                present=[]
                for label,key in items:
                    value = (views.get(key) or "").strip()
                    if value: present.append((label,value))
                if present:
                    st.markdown(f"<li><b>{heading}</b>", unsafe_allow_html=True)
                    st.markdown("<ul style='margin-top:0; margin-bottom:0;'>", unsafe_allow_html=True)
                    for label,value in present:
                        st.markdown(f"<li><i>{label}:</i> {value}</li>", unsafe_allow_html=True)
                    st.markdown("</ul></li>", unsafe_allow_html=True)
            st.markdown("</ul></li>", unsafe_allow_html=True)
        st.markdown("</ul>", unsafe_allow_html=True)
    else:
        st.info("No LCE stage activities selected.")

    # 6b) Show the LLM sections verbatim
    st.subheader("Plan (raw sections)")
    st.markdown("**Supply Chain Configuration & Action Plan**")
    st.info(parse_section(plan_text,"Supply Chain Configuration & Action Plan") or "—")
    st.markdown("**Improvement Opportunities & Risks**")
    st.info(parse_section(plan_text,"Improvement Opportunities & Risks") or "—")
    st.markdown("**Digital/AI Next Steps**")
    st.info(parse_section(plan_text,"Digital/AI Next Steps") or "—")

    # ===== Expected 5S with rationales =====
    st.markdown("**Expected 5S Maturity (with rationale)**")
    exp_raw = parse_section(plan_text, "Expected 5S Maturity")
    expected_5s = extract_expected_levels(exp_raw, five_s_levels)

    cache_key = st.session_state.get("current_cache_key")
    if "why_cache" not in st.session_state:
        st.session_state["why_cache"] = {}
    why_5s = st.session_state["why_cache"].get(cache_key) if cache_key else None
    if not why_5s:
        why_5s = explain_expected_levels(
            objective=objective,
            system_type=system_type,
            industry=industry,
            selected_stages=displayed_stages,
            five_s_levels=five_s_levels,
            expected_5s=expected_5s,
            plan_text=plan_text
        )
        if cache_key:
            st.session_state["why_cache"][cache_key] = why_5s

    rows = []
    for dim in ["Social","Sustainable","Sensing","Smart","Safe"]:
        cur = five_s_levels.get(dim, 0)
        exp = expected_5s.get(dim, cur)
        delta = exp - cur
        reasons = (why_5s.get(dim, {}) or {}).get("why", []) or []
        def _clean_reason(s: str) -> str:
            s = re.sub(r"[*_`]", "", str(s))
            s = re.sub(r"\s+", " ", s).strip()
            return s
        reasons = [_clean_reason(r) for r in reasons if r]
        rows.append({
            "Dimension": dim,
            "Current": cur,
            "Expected": exp,
            "Δ": f"{delta:+d}",
            "Why (evidence)": "; ".join(reasons[:2])
        })
    st.dataframe(pd.DataFrame(rows), use_container_width=True)

    # ================================
    # Evaluation metrics (optional)
    # ================================
    st.header("Metric Evaluation")
    st.caption(
        "Enter baseline and proposed values for evaluation metrics. "
        "Proposed values represent the expected scenario after applying the plan. "
        "Select the data source (measured, simulated, estimated, or illustrative)."
    )

    if "eval_df" not in st.session_state:
        base = pd.DataFrame(EVAL_METRICS)
        base = base.rename(columns={
            "category": "Category",
            "metric": "Metric",
            "unit": "Unit",
            "direction": "Direction",
            "driver": "Driver",
            "min": "Min",
            "max": "Max",
        })
        base["Baseline"] = None
        base["Proposed"] = None
        base["Source"] = EVAL_SOURCES[-1]
        st.session_state["eval_df"] = base

    st.caption(
        "Auto-generate illustrative metrics: baseline from current 5S levels and proposed from expected "
        "maturity, using a fixed ±5% variability. Use only for illustration unless you have "
        "measured or simulated data."
    )
    if st.button("Generate synthetic metrics"):
        synth = synthesize_metrics(five_s_levels, expected_5s, noise_pct=5.0, seed=42)
        st.session_state["eval_df"] = synth

    with st.expander("Edit metric inputs (optional)"):
        auto_prop = st.checkbox(
            "Auto-calculate proposed from baseline using expected 5S maturity",
            value=True,
        )
        display_df = st.session_state["eval_df"].drop(columns=["Driver", "Min", "Max"], errors="ignore")
        edited = st.data_editor(
            display_df,
            num_rows="fixed",
            use_container_width=True,
            column_config={
                "Baseline": st.column_config.NumberColumn("Baseline"),
                "Proposed": st.column_config.NumberColumn("Proposed"),
                "Source": st.column_config.SelectboxColumn("Source", options=EVAL_SOURCES),
            },
            disabled=["Category", "Metric", "Unit", "Direction"],
        )
        # Merge edited values back into full dataframe
        full_df = st.session_state["eval_df"].copy()
        for col in ["Baseline", "Proposed", "Source"]:
            if col in edited.columns:
                full_df[col] = edited[col]
        if auto_prop:
            full_df = compute_proposed_from_baseline(full_df, five_s_levels, expected_5s)
        st.session_state["eval_df"] = full_df

    eval_rows = compute_eval_results(st.session_state["eval_df"])
    st.session_state["eval_results"] = eval_rows
    if eval_rows:
        out_df = pd.DataFrame(eval_rows)
        st.subheader("Evaluation Summary")
        st.dataframe(out_df, use_container_width=True)
    else:
        st.info("No evaluation metrics entered yet. Use the generator or edit inputs above.")

    # 7) 5S Profiles
    st.header("5S Profiles")
    profiles_box = st.container()
    curr_df = pd.DataFrame({"Dimension": list(five_s_levels.keys()),
                            "Level": [five_s_levels[k] for k in five_s_levels]})
    fig_curr = px.line_polar(curr_df, r="Level", theta="Dimension", line_close=True, range_r=[0, 4])
    profiles_box.plotly_chart(fig_curr, use_container_width=True, key="five_s_current")
    if expected_5s:
        exp_df = pd.DataFrame({"Dimension": list(expected_5s.keys()),
                               "Level": [expected_5s[k] for k in expected_5s]})
        fig_exp = px.line_polar(exp_df, r="Level", theta="Dimension", line_close=True, range_r=[0, 4])
        profiles_box.plotly_chart(fig_exp, use_container_width=True, key="five_s_expected")

    # ================================
    # Results Q&A (chat with the plan)
    # ================================
    def _shorten(txt: str, limit: int = 2000) -> str:
        if not txt:
            return ""
        txt = str(txt).strip()
        return txt if len(txt) <= limit else (txt[:limit] + " …")

    def _build_results_context(res, five_s_levels, selected_stages, objective, system_type, industry, role):
        plan_text = res.get("plan_text","")
        sc = parse_section(plan_text, "Supply Chain Configuration & Action Plan")
        imp = parse_section(plan_text, "Improvement Opportunities & Risks")
        nxt = parse_section(plan_text, "Digital/AI Next Steps")
        exp = parse_section(plan_text, "Expected 5S Maturity")
        stage_views = res.get("stage_views", {}) or {}

        sv_lines = []
        for s in [s.split(":")[0].strip() for s in selected_stages]:
            v = stage_views.get(s, {})
            if any(v.values()):
                sv_lines.append(
                    f"- {s}: Fn={v.get('Function','')}; Org={v.get('Organization','')}; "
                    f"Info={v.get('Information','')}; Res={v.get('Resource','')}; Perf={v.get('Performance','')}"
                )

        ctx = [
            "PROJECT",
            f"- Objective: {objective}",
            f"- System Type: {system_type}",
            f"- Industry: {industry}",
            f"- Role: {role}",
            f"- Selected LCE Stages: {', '.join([s.split(':')[0] for s in selected_stages]) or 'None'}",
            f"- 5S Levels (current): " + ", ".join([f"{k}={v}" for k,v in five_s_levels.items()]),
            "",
            "[Supply Chain Configuration & Action Plan]",
            _shorten(sc, 2500),
            "",
            "[Improvement Opportunities & Risks]",
            _shorten(imp, 1200),
            "",
            "[Digital/AI Next Steps]",
            _shorten(nxt, 800),
            "",
            "[Expected 5S Maturity]",
            _shorten(exp, 400),
            "",
            "[Stage Views]",
            _shorten("\n".join(sv_lines), 1500),
            "",
        ]
        return "\n".join(ctx)

    st.header("Ask the Agent about these results")
    if "results_chat_history" not in st.session_state:
        st.session_state["results_chat_history"] = []

    # Show past conversation
    for m in st.session_state["results_chat_history"]:
        with st.chat_message(m["role"]):
            st.markdown(m["content"])

    user_q = st.chat_input("Ask about stage views, improvement logic, or suggest changes…")
    if user_q:
        st.session_state["results_chat_history"].append({"role":"user","content":user_q})
        with st.chat_message("user"):
            st.markdown(user_q)

        ctx = _build_results_context(
            res, five_s_levels, st.session_state.get("selected_stages", []),
            objective, system_type, industry, role_selected
        )

        msgs = [{"role":"system","content":
                 "You are a manufacturing systems expert. Answer strictly from the provided CONTEXT. "
                 "If something isn’t in the context, say so briefly. Be concise and actionable."}]
        msgs.append({"role":"user","content": f"CONTEXT:\n{ctx}\n\nQUESTION: {user_q}"})

        answer = llm(msgs, temperature=0.2, seed=42)
        with st.chat_message("assistant"):
            st.markdown(answer)
        st.session_state["results_chat_history"].append({"role":"assistant","content":answer})

# ------------------------ PDF export ------------------------
def to_latin1(text):
    if not isinstance(text,str): text=str(text)
    return text.encode('latin-1','replace').decode('latin-1')

def generate_pdf_report(plan_text, five_s_levels, expected_5s, why_5s, eval_rows=None):
    pdf = FPDF(); pdf.add_page(); pdf.set_font("Arial", size=12)
    pdf.cell(0,10,to_latin1("LCE + 5S Manufacturing Decision Support (Bounded AI Agent)"), ln=True, align="C")
    pdf.set_font("Arial", size=11)
    pdf.cell(0,8,to_latin1("Developed by Dr. Juana Isabel Méndez and Dr. Arturo Molina"), ln=True, align="C")
    pdf.cell(0,8,to_latin1("Date: "+datetime.now().strftime("%Y-%m-%d")), ln=True, align="C")
    pdf.ln(6)

    for head in ["Supply Chain Configuration & Action Plan","Improvement Opportunities & Risks","Digital/AI Next Steps","Expected 5S Maturity"]:
        pdf.set_font("Arial","B",12); pdf.cell(0,8,to_latin1(head), ln=True)
        pdf.set_font("Arial", size=11)
        if head == "Expected 5S Maturity":
            text = "\n".join([f"{d}: {expected_5s.get(d, five_s_levels.get(d,0))}" for d in ["Social","Sustainable","Sensing","Smart","Safe"]])
        else:
            text = parse_section(plan_text, head) or "—"
        pdf.multi_cell(0,7,to_latin1(text)); pdf.ln(2)

    if why_5s:
        pdf.set_font("Arial","B",12); pdf.cell(0,8,to_latin1("Why 5S improves (evidence-based)"), ln=True)
        pdf.set_font("Arial", size=11)
        for dim in ["Social","Sustainable","Sensing","Smart","Safe"]:
            bullets = (why_5s.get(dim, {}) or {}).get("why", []) or []
            for b in bullets[:2]:
                pdf.multi_cell(0,6,to_latin1(f"- {dim}: {b}"))

    if eval_rows:
        pdf.ln(2)
        pdf.set_font("Arial","B",12); pdf.cell(0,8,to_latin1("Evaluation Metrics (optional)"), ln=True)
        pdf.set_font("Arial", size=10)
        for r in eval_rows:
            metric = r.get("Metric","")
            unit = r.get("Unit","")
            baseline = r.get("Baseline")
            proposed = r.get("Proposed")
            delta = r.get("Δ")
            improve = r.get("Improvement (%)")
            source = r.get("Source","")
            line = f"- {metric} ({unit}): baseline={baseline}, proposed={proposed}, Δ={delta}"
            if improve is not None:
                line += f", improvement={improve:.1f}%"
            if source:
                line += f", source={source}"
            pdf.multi_cell(0,6,to_latin1(line))

    buf = BytesIO(); pdf_bytes = pdf.output(dest='S').encode('latin1')
    buf.write(pdf_bytes); buf.seek(0); return buf

res = st.session_state.get("agent_result")
if res:
    plan_text = res.get("plan_text","")
    exp_raw = parse_section(plan_text, "Expected 5S Maturity")
    expected_5s = extract_expected_levels(exp_raw, five_s_levels)
    why_5s = st.session_state.get("why_5s", {})
    eval_rows = st.session_state.get("eval_results", [])
    pdf_buf = generate_pdf_report(plan_text, five_s_levels, expected_5s, why_5s, eval_rows=eval_rows)
    st.download_button("Download Full Report (PDF)", data=pdf_buf,
                       file_name=f"{datetime.now():%Y-%m-%d}_AI_Agent_NoBOM_Report.pdf", mime="application/pdf")
